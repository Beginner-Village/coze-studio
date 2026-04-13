from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = Workbook()

HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=11)
HEADER_FILL = PatternFill("solid", fgColor="2B5797")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
CELL_FONT = Font(name="Arial", size=10)
CELL_ALIGN = Alignment(vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)

def style_header(ws, row, cols):
    for col in range(1, cols + 1):
        c = ws.cell(row=row, column=col)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = HEADER_ALIGN
        c.border = THIN_BORDER

def style_data(ws, row, cols):
    for col in range(1, cols + 1):
        c = ws.cell(row=row, column=col)
        c.font = CELL_FONT
        c.alignment = CELL_ALIGN
        c.border = THIN_BORDER

def write_table(ws, headers, data, start_row=1, col_widths=None):
    for i, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=i, value=h)
    style_header(ws, start_row, len(headers))
    for r, row_data in enumerate(data, start_row + 1):
        for c, val in enumerate(row_data, 1):
            ws.cell(row=r, column=c, value=val)
        style_data(ws, r, len(headers))
    if col_widths:
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[chr(64 + i) if i <= 26 else "A" + chr(64 + i - 26)].width = w
    return start_row + len(data) + 1

# ========== Sheet 1: 服务器分工 ==========
ws1 = wb.active
ws1.title = "服务器部署总览"

ws1.cell(row=1, column=1, value="猎鹰智能体平台 - 服务器部署总览").font = Font(bold=True, size=14, name="Arial", color="2B5797")
ws1.merge_cells("A1:G1")

headers = ["服务器IP", "主机名", "角色", "部署的容器", "对外端口", "网络模式", "备注"]
data = [
    ["30.3.165.209", "xcsitebsmgmt01", "Studio 应用服务器", "ynet-server (后端)\nynet-web (前端/Nginx)", "9888 (Web)\n8888 (API内部)\n9889 (MinIO代理)", "bridge", "智能体开发平台主入口"],
    ["30.3.165.210", "xcsitebsmqlog01", "Loop 应用服务器", "ynet-loop-app (后端)\nynet-loop-nginx (前端)", "8888 (API)\n80 (Web, 映射8082)", "host", "可观测性平台, 嵌入Studio"],
    ["30.3.165.211", "-", "Guard 应用服务器", "guard-app (Python+Nginx)", "8080 (Web+API)", "bridge", "安全围栏, 独立访问"],
]
r = write_table(ws1, headers, data, 3, [16, 18, 20, 30, 22, 12, 25])

# ========== Sheet 2: 中间件矩阵 ==========
ws2 = wb.create_sheet("中间件部署详情")

ws2.cell(row=1, column=1, value="中间件部署详情").font = Font(bold=True, size=14, name="Arial", color="2B5797")
ws2.merge_cells("A1:H1")

headers = ["中间件", "服务器IP", "端口", "版本", "用户名", "密码", "数据库/Bucket", "使用方"]
data = [
    ["OceanBase", "30.5.9.252\n(集群: 252/253/244/245)", "2883", "OB 集群", "mbank@mbank#dev", "mbankDB@123", "ai_studio / ai_loop / ai_guard", "Studio + Loop + Guard"],
    ["Redis", "30.3.165.185", "7080", "6.2.3", "-", "1q2w3e4r", "-", "Studio + Loop + Guard"],
    ["Elasticsearch", "30.3.165.196", "9200", "7.12.1", "-", "-", "project / project_draft / coze_resource / audit_logs", "Studio + Guard"],
    ["MinIO", "30.3.165.191", "9000", "2023.x (安装包)", "admin", "admin123", "openynet", "Studio + Loop"],
    ["RocketMQ", "30.3.165.182", "9876", "5.x (安装包)", "-", "-", "-", "Studio + Loop"],
    ["ClickHouse", "30.3.165.198", "9000", "安装包部署", "default", "(空)", "default", "Loop"],
    ["OneAPI (LLM代理)", "30.3.162.95", "4000", "-", "-", "sk-PiRwGz2B302b01C33uAP", "-", "Studio (模型调用)"],
]
r = write_table(ws2, headers, data, 3, [18, 22, 10, 16, 20, 22, 35, 22])

# 高亮 OneAPI 行
YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")
for col in range(1, 9):
    ws2.cell(row=10, column=col).fill = YELLOW_FILL

# ========== Sheet 3: 数据库详情 ==========
ws3 = wb.create_sheet("数据库详情")

ws3.cell(row=1, column=1, value="OceanBase 数据库详情").font = Font(bold=True, size=14, name="Arial", color="2B5797")
ws3.merge_cells("A1:F1")

headers = ["数据库名", "用途", "表数量", "连接地址", "字符集", "备注"]
data = [
    ["ai_studio", "Studio 主库", "112+", "30.5.9.252:2883", "utf8mb4", "含 space_embedding, space_rerank 表"],
    ["ai_loop", "Loop 主库", "44+", "30.5.9.252:2883", "utf8mb4", "用户: studio@ynet.com"],
    ["ai_guard", "Guard 主库", "18+", "30.5.9.252:2883", "utf8mb4", "安全策略数据"],
]
r = write_table(ws3, headers, data, 3, [16, 16, 12, 22, 12, 35])

r += 1
ws3.cell(row=r, column=1, value="DSN 连接格式").font = Font(bold=True, size=12, name="Arial", color="2B5797")
ws3.merge_cells(f"A{r}:F{r}")
r += 1

headers2 = ["使用场景", "格式", "示例"]
data2 = [
    ["MySQL 客户端", "-u 'user' -p'pass'", "mysql -h 30.5.9.252 -P 2883 -u 'mbank@mbank#dev' -p'mbankDB@123' ai_studio"],
    ["Go GORM DSN", "user:pass@tcp(host:port)/db", "mbank@mbank#dev:mbankDB@123@tcp(30.5.9.252:2883)/ai_studio?charset=utf8mb4"],
    ["Python SQLAlchemy", "URL 编码特殊字符", "mysql+asyncmy://mbank%40mbank%23dev:mbankDB%40123@30.5.9.252:2883/ai_guard"],
    ["Docker .env", "原样, # 必须加双引号", 'MYSQL_USER="mbank@mbank#dev"'],
    [".cnf 配置文件", "# 必须加双引号", 'user="mbank@mbank#dev"'],
]
write_table(ws3, headers2, data2, r, [20, 28, 60])

# ========== Sheet 4: 容器清单 ==========
ws4 = wb.create_sheet("容器清单")

ws4.cell(row=1, column=1, value="Docker 容器清单").font = Font(bold=True, size=14, name="Arial", color="2B5797")
ws4.merge_cells("A1:G1")

headers = ["容器名", "所在服务器", "镜像", "端口映射", "依赖中间件", "健康检查", "说明"]
data = [
    ["ynet-server", "209", "ynet-studio/ynet-server:latest", "8888, 9889", "OB, Redis, ES, MinIO, RMQ", "HTTP :8888", "Studio 后端 (Go/Hertz)"],
    ["ynet-web", "209", "ynet-studio/ynet-web:latest", "9888->80", "ynet-server", "HTTP :80", "Studio 前端 (Nginx+SPA)"],
    ["ynet-loop-app", "210", "ynet-loop/app:latest", "8888", "OB, Redis, CK, MinIO, RMQ", "HTTP :8888/ping", "Loop 后端 (Go)"],
    ["ynet-loop-nginx", "210", "ynet-loop/nginx:latest", "80(->8082)", "ynet-loop-app", "-", "Loop 前端 (Nginx)"],
    ["guard-app", "211", "coze-studio/guard:latest", "8080->80", "OB, Redis, ES", "HTTP /health", "Guard (Python/FastAPI+Nginx)"],
]
write_table(ws4, headers, data, 3, [18, 12, 32, 14, 28, 16, 25])

# ========== Sheet 5: 中间件使用矩阵 ==========
ws5 = wb.create_sheet("中间件使用矩阵")

ws5.cell(row=1, column=1, value="中间件 x 应用 使用矩阵").font = Font(bold=True, size=14, name="Arial", color="2B5797")
ws5.merge_cells("A1:D1")

headers = ["中间件", "Studio (209)", "Loop (210)", "Guard (211)"]
data = [
    ["OceanBase", "ai_studio (读写)", "ai_loop (读写)", "ai_guard (读写)"],
    ["Redis", "缓存, Session", "缓存", "缓存"],
    ["Elasticsearch", "项目搜索, 资源搜索, 审计日志", "-", "安全规则检索"],
    ["MinIO", "文件存储 (openynet bucket)", "日志存储", "-"],
    ["RocketMQ", "消息队列", "消息队列", "-"],
    ["ClickHouse", "-", "时序数据, Trace 存储", "-"],
    ["OneAPI (LLM)", "模型调用 (大语言模型)", "-", "-"],
]
GREEN_FILL = PatternFill("solid", fgColor="E2EFDA")
GRAY_FILL = PatternFill("solid", fgColor="F2F2F2")
r = write_table(ws5, headers, data, 3, [18, 30, 25, 25])
for row in range(4, 4 + len(data)):
    for col in range(2, 5):
        cell = ws5.cell(row=row, column=col)
        if cell.value and cell.value != "-":
            cell.fill = GREEN_FILL
        elif cell.value == "-":
            cell.fill = GRAY_FILL

# ========== Sheet 6: 访问地址 ==========
ws6 = wb.create_sheet("访问地址和账号")

ws6.cell(row=1, column=1, value="系统访问地址和账号").font = Font(bold=True, size=14, name="Arial", color="2B5797")
ws6.merge_cells("A1:E1")

headers = ["系统", "访问地址", "账号", "密码", "备注"]
data = [
    ["Studio (智能体平台)", "http://30.3.165.209:9888", "admin@ynet.com", "Admin@2026", "管理员账号"],
    ["Guard (安全围栏)", "http://30.3.165.211:8080", "admin", "admin123", "租户ID: default"],
    ["Loop (可观测性)", "嵌入 Studio 左侧菜单", "studio@ynet.com", "Ynet@2026", "通过Studio访问"],
    ["MinIO Console", "http://30.3.165.191:9000", "admin", "admin123", "对象存储管理"],
    ["OneAPI (模型代理)", "http://30.3.162.95:4000", "-", "sk-PiRwGz2B...", "LLM API 代理"],
]
write_table(ws6, headers, data, 3, [22, 32, 20, 22, 25])

# ========== Sheet 7: ES 索引 ==========
ws7 = wb.create_sheet("ES索引详情")

ws7.cell(row=1, column=1, value="Elasticsearch 索引").font = Font(bold=True, size=14, name="Arial", color="2B5797")
ws7.merge_cells("A1:D1")

headers = ["索引名", "用途", "关键字段", "使用方"]
data = [
    ["project", "项目搜索", "space_id, name, status, type, owner_id, create_time, update_time", "Studio"],
    ["project_draft", "项目草稿搜索", "space_id, name, status, type, owner_id, create_time, update_time", "Studio"],
    ["coze_resource", "资源搜索", "space_id, name, type", "Studio"],
    ["audit_logs", "审计日志", "space_id, user_id, action, timestamp", "Studio"],
]
write_table(ws7, headers, data, 3, [18, 18, 55, 12])

# ========== Sheet 8: 踩坑记录 ==========
ws8 = wb.create_sheet("踩坑记录")

ws8.cell(row=1, column=1, value="部署踩坑记录").font = Font(bold=True, size=14, name="Arial", color="2B5797")
ws8.merge_cells("A1:D1")

RED_FILL = PatternFill("solid", fgColor="FCE4EC")
headers = ["问题", "原因", "解决方案", "涉及服务器"]
data = [
    [".env 中 # 号被当注释", "mbank@mbank#dev 的 #dev 被截断", '用双引号包裹: MYSQL_USER="mbank@mbank#dev"', "209, 210"],
    ["容器无法访问 OB", "bridge 网络无法路由到外部 OB 集群", "改为 network_mode: host (Loop)", "210"],
    ["镜像地址少一段", "10.10.206 少了一个 octet", "改为 10.10.10.206", "210"],
    ["MinIO Access Key 不匹配", "凭据错误 (minioadmin->admin)", "AK=admin, SK=admin123", "209"],
    ["ES 搜索报错", "索引 mapping 为空", "删除重建索引, 添加正确的字段 mapping", "209"],
    ["MinIO 代理 9889 不通", "bridge 网络端口未映射", 'docker-compose.yml 添加 "9889:9889"', "209"],
    ["默认图标 403", "openynet bucket 没有 default_icon 文件", "用 upload-icons.sh 上传 39 个图标", "209"],
    ["智能体对话 401", "OneAPI Key 只允许特定模型", "模型名必须填 Qwen3-308-AB-Instruct-2507", "209"],
    ["Embedding/Rerank 创建无效", "前端 Form htmlType=submit 在 Modal 内静默失败", "改用 getFormApi + onClick 手动提交", "209"],
    ["space_rerank 表不存在", "Rerank 是新功能, 银行库未建表", "执行 create-tables-cdrcb3.sh 建表", "209 (OB)"],
    ["ClickHouse 端口", "安装包部署用 9000, Docker 映射 19000", "确认实际部署方式后配置端口", "210"],
    [".cnf 文件 # 注释", "MySQL 配置文件中 # 也是注释符", '值加双引号: user="mbank@mbank#dev"', "209"],
]
r = write_table(ws8, headers, data, 3, [30, 30, 42, 14])
for row in range(4, 4 + len(data)):
    for col in range(1, 5):
        ws8.cell(row=row, column=col).fill = RED_FILL

OUTPUT = "/Users/luzhipeng/Desktop/猎鹰智能体平台-部署配置表.xlsx"
wb.save(OUTPUT)
print(f"Generated: {OUTPUT}")

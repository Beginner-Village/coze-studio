# -*- coding: utf-8 -*-
"""生成 studio 全量 docx：集成教程 + 全量API(21模块) + 功能速览。"""
import os, json
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
HERE=os.path.dirname(os.path.abspath(__file__)); SD=os.path.join(HERE,'screenshots')
D=json.load(open('/tmp/studio_data.json',encoding='utf-8')); MODULES=json.load(open('/tmp/studio_modules.json',encoding='utf-8'))
BRAND=RGBColor(0x22,0x47,0xA8); INK=RGBColor(0x33,0x3A,0x45); MUTED=RGBColor(0x8A,0x93,0xA1)
doc=Document(); st=doc.styles['Normal']; st.font.name='Microsoft YaHei'; st.font.size=Pt(10.5); st.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
def cn(r,n='Microsoft YaHei'): r.font.name=n; r._element.rPr.rFonts.set(qn('w:eastAsia'),n)
def shade(o,c):
    e=OxmlElement('w:shd');e.set(qn('w:val'),'clear');e.set(qn('w:fill'),c);(o._tc.get_or_add_tcPr() if hasattr(o,'_tc') else o._p.get_or_add_pPr()).append(e)
def h1(t): p=doc.add_paragraph();r=p.add_run(t);r.bold=True;r.font.size=Pt(17);r.font.color.rgb=BRAND;cn(r)
def h2(t): p=doc.add_paragraph();r=p.add_run(t);r.bold=True;r.font.size=Pt(12.5);r.font.color.rgb=INK;cn(r)
def para(t): p=doc.add_paragraph();r=p.add_run(t);r.font.size=Pt(10.5);r.font.color.rgb=INK;cn(r)
def shot(img,cap):
    path=os.path.join(SD,img)
    if os.path.exists(path):
        doc.add_picture(path,width=Inches(6.1));doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
        c=doc.add_paragraph();c.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=c.add_run('▲ '+cap);r.italic=True;r.font.size=Pt(9);r.font.color.rgb=MUTED;cn(r)
def code(t): p=doc.add_paragraph();shade(p,'F4F2EC');r=p.add_run(t);r.font.name='Menlo';r.font.size=Pt(8.5);r.font.color.rgb=RGBColor(0x2A,0x2F,0x3A)
def tip(l,t): p=doc.add_paragraph();shade(p,'EAF2FB');r=p.add_run(l+' ');r.bold=True;r.font.color.rgb=RGBColor(0x2D,0x6B,0xD1);cn(r);r2=p.add_run(t);r2.font.size=Pt(10);cn(r2)
def warn(l,t): p=doc.add_paragraph();shade(p,'FFF7E8');r=p.add_run(l+' ');r.bold=True;r.font.color.rgb=RGBColor(0xE0,0x8A,0x1E);cn(r);r2=p.add_run(t);r2.font.size=Pt(10);cn(r2)
def table(headers,rows):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Light Grid Accent 1';t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,hd in enumerate(headers):
        c=t.rows[0].cells[i];c.text='';r=c.paragraphs[0].add_run(hd);r.bold=True;r.font.size=Pt(9);cn(r)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text='';r=cells[i].paragraphs[0].add_run(str(v));r.font.size=Pt(8.5);cn(r)
# 封面
t=doc.add_paragraph();t.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=t.add_run('猎鹰 Studio');r.bold=True;r.font.size=Pt(34);r.font.color.rgb=BRAND;cn(r)
s=doc.add_paragraph();s.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=s.add_run('全量 API · 集成手册');r.font.size=Pt(16);r.font.color.rgb=RGBColor(0x4B,0x55,0x63);cn(r)
tot=sum(len(it) for _,_,it in MODULES)
m=doc.add_paragraph();m.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=m.add_run(f'对外 API + 平台管理 API · {len(MODULES)} 模块 {tot} 接口\n生成日期：2026-06-13');r.font.size=Pt(10);r.font.color.rgb=MUTED;cn(r)
doc.add_page_break()
# 集成教程
h1('第一部分　集成教程'); h2('0. 对外开放 API 能做什么')
para('在 Studio 里搭好智能体/工作流并发布后，业务系统用 PAT 令牌通过 API 调用——对话、跑工作流、传文件。')
table(['概念','说明'],[[c[0],c[1]] for c in D['CONCEPTS']])
tip('💡 两套认证：','对外 API 用 Authorization: Bearer <PAT>；平台管理 API（知识库/插件/数据库等）走 session_key cookie。')
kinds={'h':lambda *a:h2(a[0]),'p':lambda *a:para(a[0]),'s':lambda *a:shot(a[0],a[1]),'c':lambda *a:code(a[0]),'t':lambda *a:tip(a[0],a[1]),'w':lambda *a:warn(a[0],a[1])}
for i,(tag,title,lead,items) in enumerate(D['STEPS'],1):
    doc.add_page_break();h1(f'{tag}　{title}');para(lead)
    for it in items: kinds[it[0]](*it[1:])
# 全量 API
doc.add_page_break(); h1('第二部分　API 接口（全量）')
para(f'猎鹰 Studio 共 {tot} 个接口，分 {len(MODULES)} 个功能模块。对外 API(Bearer PAT) + 平台管理 API(Cookie session)。')
for fname,auth,items in MODULES:
    doc.add_page_break(); h1(fname)
    badge='Bearer PAT' if auth=='bearer' else ('公开' if auth=='public' else 'Cookie session')
    para(f'认证：{badge}　|　接口数：{len(items)}')
    table(['方法','路径','用途'],[[m,p,n] for m,p,n,bd,d in items])
    shown=0
    for m,p,n,bd,d in items:
        if bd and shown<3:
            h2(f'{m} {n}');
            if d: para(d)
            code(json.dumps(bd,ensure_ascii=False,indent=2)); shown+=1
# 功能速览
doc.add_page_break(); h1('第三部分　管理后台功能速览')
for img,cap in [('03-project-list.png','项目开发：智能体列表'),('04-workflow-list.png','资源库：工作流'),('05-bot-edit.png','智能体编排：发布+预览调试(URL含bot_id)'),('11-card-list.png','资源库：卡片'),('10-prompt-list.png','资源库：提示词'),('09-database-list.png','资源库：数据库'),('01-pat-list.png','API授权：个人访问令牌'),('02-pat-create.png','创建访问令牌')]:
    shot(img,cap)
out=os.path.join(HERE,'猎鹰Studio-全量API集成手册.docx'); doc.save(out)
print('docx:',out); print('段落',len(doc.paragraphs),'表格',len(doc.tables),'图',len(doc.inline_shapes),'| %.1fMB'%(os.path.getsize(out)/1024/1024))

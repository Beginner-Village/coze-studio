# -*- coding: utf-8 -*-
"""生成 猎鹰 Studio 对外开放 API 集成手册 .docx（数据驱动，python-docx）。"""
import os, json
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
SD = os.path.join(HERE, 'screenshots')
D = json.load(open('/tmp/studio_data.json', encoding='utf-8'))
BRAND = RGBColor(0x22, 0x47, 0xA8); INK = RGBColor(0x33,0x3A,0x45); MUTED = RGBColor(0x8A,0x93,0xA1)

doc = Document()
st = doc.styles['Normal']; st.font.name='Microsoft YaHei'; st.font.size=Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
def cn(r,n='Microsoft YaHei'): r.font.name=n; r._element.rPr.rFonts.set(qn('w:eastAsia'),n)
def shade(o,c):
    e=OxmlElement('w:shd');e.set(qn('w:val'),'clear');e.set(qn('w:fill'),c)
    (o._tc.get_or_add_tcPr() if hasattr(o,'_tc') else o._p.get_or_add_pPr()).append(e)
def h1(t): p=doc.add_paragraph();r=p.add_run(t);r.bold=True;r.font.size=Pt(18);r.font.color.rgb=BRAND;cn(r);return p
def h2(t): p=doc.add_paragraph();r=p.add_run(t);r.bold=True;r.font.size=Pt(13.5);r.font.color.rgb=INK;cn(r);return p
def para(t): p=doc.add_paragraph();r=p.add_run(t);r.font.size=Pt(10.5);r.font.color.rgb=INK;cn(r);return p
def shot(img,cap):
    path=os.path.join(SD,img)
    if os.path.exists(path):
        doc.add_picture(path,width=Inches(6.1));doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
        c=doc.add_paragraph();c.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=c.add_run('▲ '+cap);r.italic=True;r.font.size=Pt(9);r.font.color.rgb=MUTED;cn(r)
def code(t): p=doc.add_paragraph();shade(p,'F4F2EC');r=p.add_run(t);r.font.name='Menlo';r.font.size=Pt(9);r.font.color.rgb=RGBColor(0x2A,0x2F,0x3A)
def tip(l,t): p=doc.add_paragraph();shade(p,'EAF2FB');r=p.add_run(l+' ');r.bold=True;r.font.color.rgb=RGBColor(0x2D,0x6B,0xD1);cn(r);r2=p.add_run(t);r2.font.size=Pt(10);cn(r2)
def warn(l,t): p=doc.add_paragraph();shade(p,'FFF7E8');r=p.add_run(l+' ');r.bold=True;r.font.color.rgb=RGBColor(0xE0,0x8A,0x1E);cn(r);r2=p.add_run(t);r2.font.size=Pt(10);cn(r2)
def table(headers,rows):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Light Grid Accent 1';t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,hd in enumerate(headers):
        c=t.rows[0].cells[i];c.text='';r=c.paragraphs[0].add_run(hd);r.bold=True;r.font.size=Pt(9.5);cn(r)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text='';r=cells[i].paragraphs[0].add_run(str(v));r.font.size=Pt(9);cn(r)
def endpoint(method,path):
    p=doc.add_paragraph();r=p.add_run(' '+method+'  ');r.bold=True;r.font.color.rgb=RGBColor(0xFF,0xFF,0xFF);shade(p,'2247A8')
    r2=p.add_run('  '+path);r2.font.name='Menlo';r2.bold=True;r2.font.size=Pt(11)

# 封面
t=doc.add_paragraph();t.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=t.add_run('猎鹰 Studio');r.bold=True;r.font.size=Pt(34);r.font.color.rgb=BRAND;cn(r)
s=doc.add_paragraph();s.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=s.add_run('对外开放 API · 集成手册');r.font.size=Pt(16);r.font.color.rgb=RGBColor(0x4B,0x55,0x63);cn(r)
m=doc.add_paragraph();m.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=m.add_run('PAT 令牌调用 · 智能体对话 · 工作流 · 文件上传\n生成日期：2026-06-13');r.font.size=Pt(10);r.font.color.rgb=MUTED;cn(r)
doc.add_page_break()

# 第一部分 集成教程
h1('第一部分　集成教程')
h2('0. 对外开放 API 能做什么')
para('在 Studio 里搭好智能体/工作流并发布后，业务系统用一把 PAT 令牌即可通过 API 调用它们——对话、跑工作流、传文件。')
table(['概念','说明'],[[c[0],c[1]] for c in D['CONCEPTS']])
tip('💡 集成顺序：','① 发布智能体拿 bot_id → ② 建 PAT 拿令牌 → ③ 调 /v3/chat 对话 / workflow 跑流程。')
kinds={'h':lambda *a:h2(a[0]),'p':lambda *a:para(a[0]),'s':lambda *a:shot(a[0],a[1]),'c':lambda *a:code(a[0]),'t':lambda *a:tip(a[0],a[1]),'w':lambda *a:warn(a[0],a[1])}
for i,(tag,title,lead,items) in enumerate(D['STEPS'],1):
    doc.add_page_break();h1(f'{tag}　{title}');para(lead)
    for it in items: kinds[it[0]](*it[1:])

# 第二部分 API 接口
doc.add_page_break();h1('第二部分　API 接口')
h2('BASE_URL & 认证')
para('所有对外接口用 PAT 以 Bearer 调用：Authorization: Bearer <PAT>。bot_id/conversation_id 等在 JSON 里都是字符串。SSE 接口用 curl -N。')
for method,path,name,group,desc,req,resp,curl in D['APIS']:
    doc.add_page_break();h1(name);endpoint(method,path);para(desc)
    h2('请求参数');table(['字段','类型','必填','说明'],req)
    h2('响应示例');code(resp);h2('curl');code(curl)

# 第三部分 功能速览
doc.add_page_break();h1('第三部分　管理后台功能速览')
for img,cap in D['OPS']: shot(img,cap)
# 附智能体编辑与 PAT 创建图
shot('05-bot-edit.png','智能体编排页：发布 + 预览调试，URL 含 bot_id')
shot('02-pat-create.png','创建个人访问令牌')

out=os.path.join(HERE,'猎鹰Studio-对外开放API集成手册.docx')
doc.save(out)
print('docx:',out)
print('段落:',len(doc.paragraphs),'表格:',len(doc.tables),'图:',len(doc.inline_shapes),'| %.1fMB'%(os.path.getsize(out)/1024/1024))
